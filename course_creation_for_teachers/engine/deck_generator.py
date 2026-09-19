"""
Universal Presentation Deck Generator Engine
Reads configuration from config.py and builds high-impact, branded,
strictly left-aligned presentation decks (.pptx) adhering to course standards.
"""

import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

# Ensure local engine directory is in path
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import course_creation_for_teachers.engine.config as config


def strip_legacy_logos(prs):
    """Programmatically strip legacy third-party partner logos from master layouts."""
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            for shp in list(layout.shapes):
                if shp.shape_type == 13 or "Picture" in shp.name or "Logo" in shp.name:
                    sp_elem = shp._element
                    sp_elem.getparent().remove(sp_elem)


def create_base_presentation(template_path=None):
    """Load an existing template or create a clean 16:9 widescreen presentation."""
    if template_path and os.path.exists(template_path):
        prs = Presentation(template_path)
        strip_legacy_logos(prs)
        sldIdLst = prs.slides._sldIdLst
        for sldId in list(sldIdLst):
            prs.part.drop_rel(sldId.rId)
            sldIdLst.remove(sldId)
    else:
        prs = Presentation()
        prs.slide_width = config.SLIDE_WIDTH_INCHES
        prs.slide_height = config.SLIDE_HEIGHT_INCHES
    return prs


def render_title_slide(prs, s_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0])
    for ph in list(slide.placeholders):
        sp = ph._element
        sp.getparent().remove(sp)

    # Left Hero Container
    left_hero = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(4.8), config.SLIDE_HEIGHT_INCHES)
    left_hero.fill.solid()
    left_hero.fill.fore_color.rgb = config.COLOR_DARK_NAVY
    left_hero.line.fill.background()

    # Vertical Accent Band
    gold_band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.8), 0, Inches(0.18), config.SLIDE_HEIGHT_INCHES)
    gold_band.fill.solid()
    gold_band.fill.fore_color.rgb = config.COLOR_GOLD
    gold_band.line.fill.background()

    # Left Hero Text
    tb_hero = slide.shapes.add_textbox(Inches(0.6), Inches(1.8), Inches(3.8), Inches(4.5))
    tf_hero = tb_hero.text_frame
    tf_hero.word_wrap = True

    p_inst = tf_hero.paragraphs[0]
    p_inst.text = config.INSTITUTION_ACRONYM
    p_inst.font.name = config.FONT_PRIMARY_HEADING
    p_inst.font.size = config.SIZE_TITLE_HERO
    p_inst.font.bold = True
    p_inst.font.color.rgb = config.COLOR_WHITE
    p_inst.space_after = Pt(6)

    p_tag = tf_hero.add_paragraph()
    p_tag.text = config.INSTITUTION_TAGLINE
    p_tag.font.name = config.FONT_PRIMARY_HEADING
    p_tag.font.size = Pt(20)
    p_tag.font.bold = True
    p_tag.font.color.rgb = config.COLOR_GOLD
    p_tag.space_after = Pt(20)

    p_dept = tf_hero.add_paragraph()
    p_dept.text = config.DEPARTMENT_NAME
    p_dept.font.name = config.FONT_PRIMARY_HEADING
    p_dept.font.size = Pt(22)
    p_dept.font.bold = True
    p_dept.font.color.rgb = config.COLOR_WHITE
    p_dept.space_after = Pt(10)

    p_cls = tf_hero.add_paragraph()
    p_cls.text = config.COURSE_TRACK_NAME
    p_cls.font.name = config.FONT_PRIMARY_HEADING
    p_cls.font.size = Pt(14)
    p_cls.font.color.rgb = config.COLOR_LIGHT_ICE

    # Right Content Container
    tb_right = slide.shapes.add_textbox(Inches(5.5), Inches(1.2), Inches(7.3), Inches(2.4))
    tf_right = tb_right.text_frame
    tf_right.word_wrap = True

    p_eyebrow = tf_right.paragraphs[0]
    p_eyebrow.text = s_data.get("eyebrow", config.INSTITUTION_NAME.upper())
    p_eyebrow.font.name = config.FONT_PRIMARY_HEADING
    p_eyebrow.font.size = Pt(14)
    p_eyebrow.font.bold = True
    p_eyebrow.font.color.rgb = config.COLOR_ACCENT_BLUE
    p_eyebrow.space_after = Pt(10)

    p_main = tf_right.add_paragraph()
    p_main.text = s_data.get("title", "Course Presentation")
    p_main.font.name = config.FONT_PRIMARY_HEADING
    p_main.font.size = config.SIZE_TITLE_MAIN
    p_main.font.bold = True
    p_main.font.color.rgb = config.COLOR_PRIMARY_NAVY

    # Horizontal Divider
    gold_rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(5.5), Inches(3.8), Inches(2.2), Inches(0.06))
    gold_rule.fill.solid()
    gold_rule.fill.fore_color.rgb = config.COLOR_GOLD
    gold_rule.line.fill.background()

    tb_sub = slide.shapes.add_textbox(Inches(5.5), Inches(4.0), Inches(7.3), Inches(2.5))
    tf_sub = tb_sub.text_frame
    tf_sub.word_wrap = True

    p_sub = tf_sub.paragraphs[0]
    p_sub.text = s_data.get("subtitle", "")
    p_sub.font.name = config.FONT_PRIMARY_HEADING
    p_sub.font.size = Pt(20)
    p_sub.font.bold = True
    p_sub.font.color.rgb = config.COLOR_CHARCOAL_TEXT
    p_sub.space_after = Pt(12)

    p_meta = tf_sub.add_paragraph()
    p_meta.text = s_data.get("meta", f"{config.COURSE_CODE} • {config.COREQUISITE_NAME}")
    p_meta.font.name = config.FONT_BODY
    p_meta.font.size = Pt(13)
    p_meta.font.color.rgb = config.COLOR_MUTED_TEXT

    return slide


def render_divider_slide(prs, s_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0])
    for ph in list(slide.placeholders):
        sp = ph._element
        sp.getparent().remove(sp)

    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, config.SLIDE_WIDTH_INCHES, config.SLIDE_HEIGHT_INCHES)
    bg.fill.solid()
    bg.fill.fore_color.rgb = config.COLOR_DARK_NAVY
    bg.line.fill.background()

    gold_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(2.2), Inches(0.2), Inches(2.6))
    gold_bar.fill.solid()
    gold_bar.fill.fore_color.rgb = config.COLOR_GOLD
    gold_bar.line.fill.background()

    tb_div = slide.shapes.add_textbox(Inches(1.5), Inches(2.0), Inches(10.8), Inches(3.2))
    tf_div = tb_div.text_frame
    tf_div.word_wrap = True

    p_num = tf_div.paragraphs[0]
    p_num.text = s_data.get("sec_num", "SECTION")
    p_num.font.name = config.FONT_PRIMARY_HEADING
    p_num.font.size = Pt(14)
    p_num.font.bold = True
    p_num.font.color.rgb = config.COLOR_GOLD
    p_num.space_after = Pt(8)

    p_dt = tf_div.add_paragraph()
    p_dt.text = s_data.get("title", "Section Title")
    p_dt.font.name = config.FONT_PRIMARY_HEADING
    p_dt.font.size = config.SIZE_SECTION_TITLE
    p_dt.font.bold = True
    p_dt.font.color.rgb = config.COLOR_WHITE
    p_dt.space_after = Pt(10)

    p_ds = tf_div.add_paragraph()
    p_ds.text = s_data.get("subtitle", "")
    p_ds.font.name = config.FONT_BODY
    p_ds.font.size = Pt(18)
    p_ds.font.color.rgb = config.COLOR_LIGHT_ICE

    return slide


def add_slide_header_and_footer(slide, title_text, eyebrow_text=None):
    tb_head = slide.shapes.add_textbox(config.MARGIN_LEFT_STANDARD, config.MARGIN_TOP_STANDARD, config.HEADER_WIDTH, Inches(1.3))
    tf_head = tb_head.text_frame
    tf_head.word_wrap = True

    p_eyebrow = tf_head.paragraphs[0]
    p_eyebrow.text = eyebrow_text if eyebrow_text else f"{config.INSTITUTION_ACRONYM} {config.COURSE_CODE} • INSTRUCTIONAL MODULE"
    p_eyebrow.font.name = config.FONT_PRIMARY_HEADING
    p_eyebrow.font.size = config.SIZE_EYEBROW
    p_eyebrow.font.bold = True
    p_eyebrow.font.color.rgb = config.COLOR_WARM_GOLD
    p_eyebrow.space_after = Pt(4)

    p_title = tf_head.add_paragraph()
    p_title.text = title_text
    p_title.font.name = config.FONT_PRIMARY_HEADING
    p_title.font.size = config.SIZE_SLIDE_HEADING
    p_title.font.bold = True
    p_title.font.color.rgb = config.COLOR_PRIMARY_NAVY

    tb_foot = slide.shapes.add_textbox(config.MARGIN_LEFT_STANDARD, Inches(6.9), config.HEADER_WIDTH, Inches(0.4))
    p_foot = tb_foot.text_frame.paragraphs[0]
    p_foot.text = config.FOOTER_TEXT
    p_foot.font.name = config.FONT_PRIMARY_HEADING
    p_foot.font.size = config.SIZE_FOOTER
    p_foot.font.bold = True
    p_foot.font.color.rgb = config.COLOR_PRIMARY_NAVY


def render_two_column_slide(prs, s_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0])
    for ph in list(slide.placeholders):
        sp = ph._element
        sp.getparent().remove(sp)

    add_slide_header_and_footer(slide, s_data.get("title", ""))

    # Left Card
    left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), config.CARD_TOP_STANDARD, Inches(5.6), config.CARD_HEIGHT_STANDARD)
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = config.CARD_BG_PRIMARY
    left_card.line.color.rgb = config.CARD_LINE_PRIMARY
    left_card.line.width = Pt(1.5)

    tf_l = left_card.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Inches(0.25)
    tf_l.margin_top = Inches(0.25)

    p_lh = tf_l.paragraphs[0]
    p_lh.text = s_data.get("left_header", "Dimension A")
    p_lh.font.name = config.FONT_PRIMARY_HEADING
    p_lh.font.size = config.SIZE_CARD_HEADER
    p_lh.font.bold = True
    p_lh.font.color.rgb = config.COLOR_PRIMARY_NAVY
    p_lh.alignment = PP_ALIGN.LEFT
    p_lh.space_after = Pt(14)

    for anchor, body in s_data.get("left_items", []):
        p = tf_l.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(12)
        r1 = p.add_run()
        r1.text = anchor
        r1.font.bold = True
        r1.font.name = config.FONT_PRIMARY_HEADING
        r1.font.size = Pt(14)
        r1.font.color.rgb = config.COLOR_PRIMARY_NAVY
        r2 = p.add_run()
        r2.text = body
        r2.font.name = config.FONT_BODY
        r2.font.size = config.SIZE_BODY_TEXT
        r2.font.color.rgb = config.COLOR_CHARCOAL_TEXT

    # Right Card
    right_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), config.CARD_TOP_STANDARD, Inches(5.6), config.CARD_HEIGHT_STANDARD)
    right_card.fill.solid()
    right_card.fill.fore_color.rgb = config.CARD_BG_SECONDARY
    right_card.line.color.rgb = config.CARD_LINE_SECONDARY
    right_card.line.width = Pt(1.5)

    tf_r = right_card.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = Inches(0.25)
    tf_r.margin_top = Inches(0.25)

    p_rh = tf_r.paragraphs[0]
    p_rh.text = s_data.get("right_header", "Dimension B")
    p_rh.font.name = config.FONT_PRIMARY_HEADING
    p_rh.font.size = config.SIZE_CARD_HEADER
    p_rh.font.bold = True
    p_rh.font.color.rgb = config.COLOR_PRIMARY_NAVY
    p_rh.alignment = PP_ALIGN.LEFT
    p_rh.space_after = Pt(14)

    for anchor, body in s_data.get("right_items", []):
        p = tf_r.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(12)
        r1 = p.add_run()
        r1.text = anchor
        r1.font.bold = True
        r1.font.name = config.FONT_PRIMARY_HEADING
        r1.font.size = Pt(14)
        r1.font.color.rgb = config.COLOR_PRIMARY_NAVY
        r2 = p.add_run()
        r2.text = body
        r2.font.name = config.FONT_BODY
        r2.font.size = config.SIZE_BODY_TEXT
        r2.font.color.rgb = config.COLOR_CHARCOAL_TEXT

    return slide


def render_normal_slide(prs, s_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0])
    for ph in list(slide.placeholders):
        sp = ph._element
        sp.getparent().remove(sp)

    add_slide_header_and_footer(slide, s_data.get("title", ""))

    content_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), config.CARD_TOP_STANDARD, config.HEADER_WIDTH, config.CARD_HEIGHT_STANDARD)
    content_card.fill.solid()
    content_card.fill.fore_color.rgb = config.CARD_BG_PRIMARY
    content_card.line.color.rgb = config.CARD_LINE_PRIMARY
    content_card.line.width = Pt(1.5)

    tf_c = content_card.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = Inches(0.4)
    tf_c.margin_top = Inches(0.35)

    items = s_data.get("items", [])
    for i_idx, (anchor, body) in enumerate(items):
        p = tf_c.paragraphs[0] if i_idx == 0 else tf_c.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(14)
        r1 = p.add_run()
        r1.text = anchor
        r1.font.bold = True
        r1.font.name = config.FONT_PRIMARY_HEADING
        r1.font.size = Pt(15)
        r1.font.color.rgb = config.COLOR_PRIMARY_NAVY
        r2 = p.add_run()
        r2.text = body
        r2.font.name = config.FONT_BODY
        r2.font.size = Pt(14.5)
        r2.font.color.rgb = config.COLOR_CHARCOAL_TEXT

    return slide


def render_code_slide(prs, s_data):
    slide = prs.slides.add_slide(prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0])
    for ph in list(slide.placeholders):
        sp = ph._element
        sp.getparent().remove(sp)

    add_slide_header_and_footer(slide, s_data.get("title", ""))

    # Left Spec Card
    left_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), config.CARD_TOP_STANDARD, Inches(5.2), config.CARD_HEIGHT_STANDARD)
    left_card.fill.solid()
    left_card.fill.fore_color.rgb = config.CARD_BG_PRIMARY
    left_card.line.color.rgb = config.CARD_LINE_PRIMARY
    left_card.line.width = Pt(1.5)

    tf_l = left_card.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Inches(0.22)
    tf_l.margin_top = Inches(0.22)

    p_lh = tf_l.paragraphs[0]
    p_lh.text = s_data.get("left_header", "Concept & Specification")
    p_lh.font.name = config.FONT_PRIMARY_HEADING
    p_lh.font.size = Pt(20)
    p_lh.font.bold = True
    p_lh.font.color.rgb = config.COLOR_PRIMARY_NAVY
    p_lh.alignment = PP_ALIGN.LEFT
    p_lh.space_after = Pt(12)

    for anchor, body in s_data.get("left_points", []):
        p = tf_l.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        r1 = p.add_run()
        r1.text = anchor
        r1.font.bold = True
        r1.font.name = config.FONT_PRIMARY_HEADING
        r1.font.size = Pt(13.5)
        r1.font.color.rgb = config.COLOR_PRIMARY_NAVY
        r2 = p.add_run()
        r2.text = body
        r2.font.name = config.FONT_BODY
        r2.font.size = Pt(13)
        r2.font.color.rgb = config.COLOR_CHARCOAL_TEXT

    # Right Code Container (STRICTLY LEFT-ALIGNED)
    code_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.3), config.CARD_TOP_STANDARD, Inches(6.2), config.CARD_HEIGHT_STANDARD)
    code_card.fill.solid()
    code_card.fill.fore_color.rgb = config.CODE_CONTAINER_BG
    code_card.line.color.rgb = config.CODE_CONTAINER_LINE
    code_card.line.width = Pt(1.5)

    tf_c = code_card.text_frame
    tf_c.word_wrap = True
    tf_c.margin_left = Inches(0.25)
    tf_c.margin_top = Inches(0.25)

    code_text = s_data.get("code_snippet", "# Code Container").strip()
    lines = code_text.split("\n")
    for l_idx, line in enumerate(lines):
        p = tf_c.paragraphs[0] if l_idx == 0 else tf_c.add_paragraph()
        p.text = line if line else " "
        p.font.name = config.FONT_CODE_MONO
        p.font.size = config.SIZE_CODE_TEXT
        p.font.color.rgb = config.CODE_SYNTAX_TEXT
        p.alignment = config.CODE_ALIGNMENT        # STRICT LEFT ALIGNMENT
        p.space_after = config.CODE_LINE_SPACING

    return slide


def render_cards_slide(prs, s_data):
    """Render 3 cards side-by-side in a horizontal grid."""
    slide = prs.slides.add_slide(prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0])
    for ph in list(slide.placeholders):
        sp = ph._element
        sp.getparent().remove(sp)

    add_slide_header_and_footer(slide, s_data.get("title", ""))

    cards_list = s_data.get("cards", [])
    card_width = Inches(3.6)
    card_gap = Inches(0.4)
    start_left = Inches(0.8)

    for c_idx, card_info in enumerate(cards_list[:3]):
        left_pos = start_left + c_idx * (card_width + card_gap)
        c_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, config.CARD_TOP_STANDARD, card_width, config.CARD_HEIGHT_STANDARD)
        c_shape.fill.solid()
        c_shape.fill.fore_color.rgb = config.CARD_BG_PRIMARY
        c_shape.line.color.rgb = config.CARD_LINE_PRIMARY
        c_shape.line.width = Pt(1.5)

        tf = c_shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_top = Inches(0.25)
        tf.margin_right = Inches(0.2)

        p_h = tf.paragraphs[0]
        p_h.text = card_info.get("title", f"Card {c_idx+1}")
        p_h.font.name = config.FONT_PRIMARY_HEADING
        p_h.font.size = Pt(17)
        p_h.font.bold = True
        p_h.font.color.rgb = config.COLOR_PRIMARY_NAVY
        p_h.alignment = PP_ALIGN.LEFT
        p_h.space_after = Pt(12)

        items = card_info.get("items", [])
        for item in items:
            p = tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            p.space_after = Pt(8)
            r = p.add_run()
            if isinstance(item, tuple) or isinstance(item, list):
                r.text = f"{item[0]}: {item[1]}"
            else:
                r.text = str(item)
            r.font.name = config.FONT_BODY
            r.font.size = Pt(13)
            r.font.color.rgb = config.COLOR_CHARCOAL_TEXT

    return slide


def render_warning_slide(prs, s_data):
    """Render a warning or high-stakes caution slide with alerting visuals."""
    from pptx.dml.color import RGBColor
    slide = prs.slides.add_slide(prs.slide_layouts[6] if len(prs.slide_layouts) > 6 else prs.slide_layouts[0])
    for ph in list(slide.placeholders):
        sp = ph._element
        sp.getparent().remove(sp)

    add_slide_header_and_footer(slide, s_data.get("title", ""))

    warn_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), config.CARD_TOP_STANDARD, config.HEADER_WIDTH, config.CARD_HEIGHT_STANDARD)
    warn_card.fill.solid()
    warn_card.fill.fore_color.rgb = RGBColor(254, 242, 242)  # Light rose
    warn_card.line.color.rgb = RGBColor(220, 38, 38)        # Crimson red
    warn_card.line.width = Pt(2.0)

    tf = warn_card.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.4)
    tf.margin_top = Inches(0.35)
    tf.margin_right = Inches(0.4)

    p_wh = tf.paragraphs[0]
    p_wh.text = s_data.get("card_title", "CRITICAL WARNING / COMMON PITFALL")
    p_wh.font.name = config.FONT_PRIMARY_HEADING
    p_wh.font.size = Pt(20)
    p_wh.font.bold = True
    p_wh.font.color.rgb = RGBColor(185, 28, 28)
    p_wh.alignment = PP_ALIGN.LEFT
    p_wh.space_after = Pt(16)

    warning_items = s_data.get("warning_items", s_data.get("items", []))
    for item in warning_items:
        p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(12)
        r = p.add_run()
        if isinstance(item, tuple) or isinstance(item, list):
            r.text = f"[!] {item[0]}: {item[1]}"
        else:
            r.text = f"[!] {item}"
        r.font.name = config.FONT_BODY
        r.font.size = Pt(14)
        r.font.color.rgb = config.COLOR_CHARCOAL_TEXT

    return slide


def build_deck(slides_data, output_path, template_path=None):
    """Builds and saves a presentation from a slides_data list."""
    prs = create_base_presentation(template_path)
    print(f"Building presentation with {len(slides_data)} slides...")

    for idx, s_data in enumerate(slides_data):
        stype = s_data.get("archetype", s_data.get("type", "normal")).lower()
        notes_text = s_data.get("notes", "")

        # Normalize archetype aliases
        if stype in ["title", "hero"]:
            slide = render_title_slide(prs, s_data)
        elif stype in ["divider", "section"]:
            slide = render_divider_slide(prs, s_data)
        elif stype in ["two_column", "split_concept", "split"]:
            # Normalize split concept keys if needed
            if "headline_left" in s_data and "left_header" not in s_data:
                s_data["left_header"] = s_data["headline_left"]
            if "headline_right" in s_data and "right_header" not in s_data:
                s_data["right_header"] = s_data["headline_right"]
            if "body_left" in s_data and "left_items" not in s_data:
                s_data["left_items"] = [("", b) if isinstance(b, str) else b for b in s_data["body_left"]]
            if "body_right" in s_data and "right_items" not in s_data:
                s_data["right_items"] = [("", b) if isinstance(b, str) else b for b in s_data["body_right"]]
            slide = render_two_column_slide(prs, s_data)
        elif stype in ["cards", "grid", "three_column"]:
            slide = render_cards_slide(prs, s_data)
        elif stype in ["warning", "caution", "alert"]:
            slide = render_warning_slide(prs, s_data)
        elif stype in ["code", "lab", "implementation"]:
            if "code" in s_data and "code_snippet" not in s_data:
                s_data["code_snippet"] = s_data["code"]
            slide = render_code_slide(prs, s_data)
        else:
            slide = render_normal_slide(prs, s_data)

        if notes_text:
            slide.notes_slide.notes_text_frame.text = notes_text

        print(f"  Slide {idx+1:02d} [{stype.upper()}]: {s_data.get('title', '')[:45]}...")

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    prs.save(output_path)
    size_mb = os.path.getsize(output_path) / (1024 * 1024)
    print(f"[PASS] Saved presentation to '{output_path}' ({size_mb:.2f} MB, {len(prs.slides)} slides)")
    return output_path


class UniversalDeckGenerator:
    """Class wrapper around build_deck for OOP-style invocations."""
    def __init__(self, template_path=None):
        self.template_path = template_path

    def generate(self, slides_data, output_path):
        return build_deck(slides_data, output_path, self.template_path)

